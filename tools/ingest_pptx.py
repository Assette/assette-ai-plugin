"""Phase 1 — Ingest & Element Inventory.

Parse one or more .pptx files into a normalized JSON representation that every
downstream phase reads from. Pure deterministic I/O — no LLM calls. Underinvesting
here compounds; the JSON shape produced by this tool is the working-state contract
for the entire pipeline.

See `docs/APPROACH.md` § Phase 1 and `workspace/README.md` for context.

================================================================================
OUTPUT SCHEMA — workspace/corpus_inventory.json
================================================================================
{
  "schema_version": "0.1.0",
  "ingested_at": "2026-05-18T12:34:56Z",
  "corpus": {
    "input_path": "<absolute path to input directory>",
    "deck_count": <int>,
    "element_count": <int>
  },
  "decks": [
    {
      "deck_id": "<sha256[:12] of file bytes>",
      "filename": "<basename>",
      "relative_path": "<path relative to corpus root>",
      "metadata": {
        "slide_width_emu": <int>,
        "slide_height_emu": <int>,
        "orientation": "landscape" | "portrait",
        "slide_count": <int>,
        "theme_colors": [<hex>, ...],
        "fonts": {"major": <name>, "minor": <name>}
      },
      "slides": [
        {
          "slide_index": <int, 0-based>,
          "slide_id": "<deck_id>:<slide_index>",
          "layout_name": <str>,
          "elements": [
            {
              "element_id": "<slide_id>:<element_index>",
              "element_index": <int, 0-based>,
              "kind": "text" | "table" | "chart" | "picture" | "placeholder" | "group" | "shape",
              "position": {"left_emu": <int>, "top_emu": <int>, "width_emu": <int>, "height_emu": <int>},
              "name": <str | null>,
              "content": { ... kind-specific payload ... },
              "raw": { ... selected raw properties for round-tripping ... }
            }
          ]
        }
      ]
    }
  ]
}

KIND-SPECIFIC CONTENT
- text:        {"text": <str>, "runs": [{"text": <str>, "bold": <bool>, "italic": <bool>, "size_pt": <float|null>, "font_name": <str|null>, "color_hex": <str|null>}]}
- table:       {"rows": <int>, "cols": <int>, "cells": [[{"text": <str>, "number_format": <str|null>}]], "has_header": <bool>}
- chart:       {"chart_type": <str>, "categories": [<str>], "series": [{"name": <str>, "values": [<float|str>]}], "embedded_xlsx_present": <bool>}
- picture:     {"image_format": <str>, "alt_text": <str|null>, "sha256": <str>}
- placeholder: {"placeholder_type": <str>, "idx": <int>}
- group:       {"member_count": <int>, "member_ids": [<element_id>]}
- shape:       {"shape_type": <str>, "text": <str|null>}
================================================================================
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from _inventory_common import (
    SCHEMA_VERSION,
    build_corpus,
    sha256_bytes as _sha256_bytes,
    sha256_file as _sha256_file,
)

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as exc:  # pragma: no cover
    print(
        "ingest_pptx.py requires python-pptx. Install with: pip install -r requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1) from exc


# SCHEMA_VERSION is re-exported from _inventory_common so downstream code and tests
# that read ingest_pptx.SCHEMA_VERSION continue to resolve a single source of truth.
SOURCE_FORMAT = "pptx"


# --------------------------------------------------------------------------- Shape extraction


def _emu_position(shape) -> dict[str, int]:
    return {
        "left_emu": int(shape.left or 0),
        "top_emu": int(shape.top or 0),
        "width_emu": int(shape.width or 0),
        "height_emu": int(shape.height or 0),
    }


def _extract_text(shape) -> dict[str, Any]:
    """Text frame with run-level styling."""
    if not shape.has_text_frame:
        return {"text": "", "runs": []}
    runs: list[dict[str, Any]] = []
    full_text_parts: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            color_hex = None
            try:
                if run.font.color and run.font.color.rgb is not None:
                    color_hex = f"#{run.font.color.rgb}"
            except (AttributeError, ValueError):
                color_hex = None
            runs.append(
                {
                    "text": run.text,
                    "bold": bool(run.font.bold),
                    "italic": bool(run.font.italic),
                    "size_pt": float(run.font.size.pt) if run.font.size else None,
                    "font_name": run.font.name,
                    "color_hex": color_hex,
                }
            )
            full_text_parts.append(run.text)
        full_text_parts.append("\n")
    return {"text": "".join(full_text_parts).rstrip("\n"), "runs": runs}


def _extract_table(shape) -> dict[str, Any]:
    table = shape.table
    rows = len(table.rows)
    cols = len(table.columns)
    cells: list[list[dict[str, Any]]] = []
    for row in table.rows:
        row_out: list[dict[str, Any]] = []
        for cell in row.cells:
            row_out.append(
                {
                    "text": cell.text.strip(),
                    "number_format": None,  # python-pptx does not expose this directly; left null
                }
            )
        cells.append(row_out)
    has_header = bool(table.first_row) if hasattr(table, "first_row") else False
    return {"rows": rows, "cols": cols, "cells": cells, "has_header": has_header}


def _extract_chart(shape) -> dict[str, Any]:
    chart = shape.chart
    chart_type = str(chart.chart_type) if chart.chart_type is not None else "unknown"
    categories: list[str] = []
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
    except (IndexError, AttributeError, ValueError):
        categories = []
    series: list[dict[str, Any]] = []
    try:
        for ser in chart.series:
            values = []
            try:
                values = [float(v) if v is not None else None for v in ser.values]
            except (TypeError, ValueError):
                values = list(ser.values)
            series.append({"name": ser.name, "values": values})
    except Exception:  # noqa: BLE001 — python-pptx chart parsing is fragile
        series = []
    # The embedded xlsx is at chart.part.chart_workbook.xlsx_part — check presence only
    embedded_present = False
    try:
        embedded_present = chart.part.chart_workbook.xlsx_part is not None
    except AttributeError:
        embedded_present = False
    return {
        "chart_type": chart_type,
        "categories": categories,
        "series": series,
        "embedded_xlsx_present": embedded_present,
    }


def _extract_picture(shape) -> dict[str, Any]:
    image = shape.image
    blob = image.blob
    return {
        "image_format": image.ext,
        "alt_text": getattr(shape, "name", None),
        "sha256": _sha256_bytes(blob),
    }


def _extract_placeholder(shape) -> dict[str, Any]:
    ph = shape.placeholder_format
    return {
        "placeholder_type": str(ph.type) if ph and ph.type else "unknown",
        "idx": ph.idx if ph else -1,
    }


# --------------------------------------------------------------------------- Walk


@dataclass
class _ElementWalker:
    deck_id: str
    slide_index: int
    element_counter: int = 0
    elements: list[dict[str, Any]] = field(default_factory=list)

    def _slide_id(self) -> str:
        return f"{self.deck_id}:{self.slide_index}"

    def _next_id(self) -> tuple[int, str]:
        idx = self.element_counter
        self.element_counter += 1
        return idx, f"{self._slide_id()}:{idx}"

    def walk_shape(self, shape, *, parent_member_ids: list[str] | None = None) -> str:
        idx, eid = self._next_id()
        kind, content = self._classify_and_extract(shape, idx, eid)
        record = {
            "element_id": eid,
            "element_index": idx,
            "kind": kind,
            "position": _emu_position(shape),
            "name": getattr(shape, "name", None),
            "content": content,
            "raw": {"shape_id": int(getattr(shape, "shape_id", -1))},
        }
        self.elements.append(record)
        if parent_member_ids is not None:
            parent_member_ids.append(eid)
        return eid

    def _classify_and_extract(self, shape, idx: int, eid: str) -> tuple[str, dict[str, Any]]:
        # Groups expand into members; emit the group record AND recurse.
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            member_ids: list[str] = []
            for child in shape.shapes:
                self.walk_shape(child, parent_member_ids=member_ids)
            return "group", {"member_count": len(member_ids), "member_ids": member_ids}
        if shape.has_table:
            return "table", _extract_table(shape)
        if shape.has_chart:
            return "chart", _extract_chart(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return "picture", _extract_picture(shape)
        if shape.is_placeholder:
            payload: dict[str, Any] = _extract_placeholder(shape)
            # Placeholders may also have text — fold it in.
            if shape.has_text_frame:
                payload["text"] = _extract_text(shape)
            return "placeholder", payload
        if shape.has_text_frame:
            return "text", _extract_text(shape)
        return "shape", {
            "shape_type": str(shape.shape_type) if shape.shape_type is not None else "unknown",
            "text": None,
        }


# --------------------------------------------------------------------------- Deck


def _deck_metadata(prs: Presentation) -> dict[str, Any]:
    width = int(prs.slide_width or 0)
    height = int(prs.slide_height or 0)
    orientation = "landscape" if width >= height else "portrait"
    # Theme colors and fonts: best-effort
    theme_colors: list[str] = []
    fonts = {"major": None, "minor": None}
    try:
        theme = prs.slide_master.element.getparent()  # noqa: SLF001
        # python-pptx does not expose theme dicts directly; leaving as best-effort.
        # Downstream consumers should not depend on theme_colors/fonts at v0.1.
        _ = theme
    except Exception:  # noqa: BLE001
        pass
    return {
        "slide_width_emu": width,
        "slide_height_emu": height,
        "orientation": orientation,
        "slide_count": len(prs.slides),
        "theme_colors": theme_colors,
        "fonts": fonts,
    }


def ingest_deck(path: Path, relative_to: Path) -> dict[str, Any]:
    deck_id = _sha256_file(path)[:12]
    prs = Presentation(str(path))
    slides_out: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(prs.slides):
        walker = _ElementWalker(deck_id=deck_id, slide_index=slide_index)
        for shape in slide.shapes:
            walker.walk_shape(shape)
        slides_out.append(
            {
                "slide_index": slide_index,
                "slide_id": f"{deck_id}:{slide_index}",
                "page_kind": "slide",
                "layout_name": slide.slide_layout.name,
                "elements": walker.elements,
            }
        )
    return {
        "deck_id": deck_id,
        "filename": path.name,
        "relative_path": str(path.relative_to(relative_to)).replace("\\", "/"),
        "source_format": SOURCE_FORMAT,
        "metadata": _deck_metadata(prs),
        "slides": slides_out,
    }


def ingest_corpus(input_path: Path) -> dict[str, Any]:
    input_path = input_path.resolve()
    if input_path.is_file():
        deck_paths = [input_path]
        relative_to = input_path.parent
    else:
        deck_paths = sorted(input_path.rglob("*.pptx"))
        relative_to = input_path
    decks = [ingest_deck(p, relative_to) for p in deck_paths]
    return build_corpus(decks, input_path)


# --------------------------------------------------------------------------- CLI


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Ingest PPTX corpus into normalized JSON.")
    parser.add_argument("--input", required=True, help="Path to a .pptx file or a directory of decks.")
    parser.add_argument(
        "--output",
        required=True,
        help="Path to write corpus_inventory.json (parent dirs will not be auto-created).",
    )
    args = parser.parse_args(argv)
    input_path = Path(args.input)
    output_path = Path(args.output)
    if not input_path.exists():
        print(f"Input path does not exist: {input_path}", file=sys.stderr)
        return 2
    inventory = ingest_corpus(input_path)
    output_path.write_text(json.dumps(inventory, indent=2), encoding="utf-8")
    print(
        f"Ingested {inventory['corpus']['deck_count']} deck(s), "
        f"{inventory['corpus']['element_count']} element(s) -> {output_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
